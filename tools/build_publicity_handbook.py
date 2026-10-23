"""Build the publicity DOCX from its reviewed UTF-8 Markdown sources.

Run with the Codex packaged Python, rather than changing project dependencies.
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE_DIR = ROOT / "docs" / "publicity"


def set_fonts(run, size=None, bold=None, code=False):
    run.font.name = "Consolas" if code else "Calibri"
    run.font.color.rgb = RGBColor(0, 0, 0)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")


def append_link(paragraph, label, href):
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), paragraph.part.relate_to(
        href, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True
    ))
    run = OxmlElement("w:r")
    prop = OxmlElement("w:rPr")
    font = OxmlElement("w:rFonts")
    font.set(qn("w:ascii"), "Calibri")
    font.set(qn("w:hAnsi"), "Calibri")
    font.set(qn("w:eastAsia"), "Microsoft YaHei")
    prop.append(font)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "16476A")
    prop.append(color)
    run.append(prop)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^\)]+\))")


def add_inline(paragraph, text, size=None):
    for token in INLINE.split(text):
        if not token:
            continue
        if token.startswith("[") and "](" in token:
            match = re.fullmatch(r"\[([^\]]+)\]\(([^\)]+)\)", token)
            label, href = match.groups()
            if href.startswith(("https://", "http://")):
                append_link(paragraph, label, href)
            else:
                set_fonts(paragraph.add_run(f"{label}（项目资料 {href}）"), size=size)
        elif token.startswith("**"):
            set_fonts(paragraph.add_run(token[2:-2]), size=size, bold=True)
        elif token.startswith("`"):
            set_fonts(paragraph.add_run(token[1:-1]), size=size, code=True)
        else:
            set_fonts(paragraph.add_run(token), size=size)


def table(document, records):
    cols = len(records[0])
    widths = {3: [2.35, 1.62, 2.53], 4: [1.42, 1.26, 1.72, 2.10]}[cols]
    if "总输出吞吐" in records[0]:
        widths = [3.10, 1.65, 1.75] if cols == 3 else [2.60, 1.02, 1.52, 1.36]
    result = document.add_table(rows=1, cols=cols)
    result.alignment = WD_TABLE_ALIGNMENT.CENTER
    result.autofit = False
    properties = result._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        edge = OxmlElement(f"w:{side}")
        for key, value in [("val", "single"), ("sz", "4"), ("color", "D9D9D9")]:
            edge.set(qn(f"w:{key}"), value)
        borders.append(edge)
    properties.append(borders)
    margin = OxmlElement("w:tblCellMar")
    for side, value in [("top", "90"), ("bottom", "90"), ("left", "95"), ("right", "95")]:
        edge = OxmlElement(f"w:{side}")
        edge.set(qn("w:w"), value)
        edge.set(qn("w:type"), "dxa")
        margin.append(edge)
    properties.append(margin)
    for column, width in zip(result.columns, widths):
        column.width = Inches(width)
    for i, record in enumerate(records):
        row = result.rows[0] if i == 0 else result.add_row()
        if i == 0:
            header = OxmlElement("w:tblHeader")
            row._tr.get_or_add_trPr().append(header)
        no_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(no_split)
        for j, (cell, value) in enumerate(zip(row.cells, record)):
            cell.width = Inches(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            if "吞吐" in records[0][j] or "P95" in records[0][j] or records[0][j] == "精度":
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(p, value, size=10.5)
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "173D55" if i == 0 else ("F1F5F8" if i % 2 == 0 else "FFFFFF"))
            cell._tc.get_or_add_tcPr().append(shading)
            if i == 0:
                for run in p.runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
    after = document.add_paragraph()
    after.paragraph_format.space_after = Pt(2)
    after.paragraph_format.space_before = Pt(0)
    after.paragraph_format.line_spacing = 0.35


def add_markdown(document, path, *, compact=False):
    starting_paragraph = len(document.paragraphs)
    lines = path.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith("```"):
            code = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            # One paragraph per logical code line; allow safe wrapping without
            # turning the code into a clipped table or shaded callout box.
            for k, value in enumerate(code):
                p = document.add_paragraph(style="Code")
                p.paragraph_format.keep_with_next = k < len(code) - 1
                if k == len(code) - 1:
                    p.paragraph_format.space_after = Pt(7)
                set_fonts(p.add_run(value or " "), size=9.5, code=True)
            i += 1
            continue
        if line.startswith("|"):
            records = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", c) for c in cells):
                    records.append(cells)
                i += 1
            table(document, records)
            continue
        if line.startswith("#"):
            hashes = len(line) - len(line.lstrip("#"))
            text = line[hashes:].strip()
            # The first Markdown title becomes a section within the compilation.
            level = min(hashes, 3)
            p = document.add_paragraph(style=f"Heading {level}")
            add_inline(p, text)
            i += 1
            continue
        if line.startswith("- "):
            p = document.add_paragraph(style="List Bullet")
            add_inline(p, line[2:])
        else:
            p = document.add_paragraph()
            add_inline(p, line)
            if line.startswith("vLLM 本身已有简洁"):
                p.paragraph_format.keep_together = True
        i += 1
    if compact:
        for p in document.paragraphs[starting_paragraph:]:
            p.paragraph_format.line_spacing = 1.18
            if p.style.name == "Heading 2":
                p.paragraph_format.space_before = Pt(8)
                p.paragraph_format.space_after = Pt(5)
                for run in p.runs:
                    run.font.size = Pt(12)
            elif p.style.name != "Heading 1":
                p.paragraph_format.space_after = Pt(5)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, default=SOURCE_DIR / "SaddleLLM_产品手册与公众号宣传文案.docx")
    args = parser.parse_args()
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = section.right_margin = Inches(1)
    section.top_margin = Inches(0.80)
    section.bottom_margin = Inches(0.78)
    for name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "List Bullet"]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.widow_control = True
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.30
    normal.paragraph_format.space_after = Pt(7)
    for name, size in [("Title", 22), ("Heading 1", 17), ("Heading 2", 14), ("Heading 3", 12)]:
        doc.styles[name].font.size = Pt(size)
        doc.styles[name].paragraph_format.space_before = Pt(13 if name != "Title" else 0)
        doc.styles[name].paragraph_format.space_after = Pt(7)
        doc.styles[name].paragraph_format.keep_with_next = True
    code = doc.styles.add_style("Code", 1)
    code.font.name = "Consolas"
    code.font.size = Pt(9.5)
    code.paragraph_format.line_spacing = 1.1
    code.paragraph_format.space_after = Pt(1)
    code.paragraph_format.left_indent = Inches(0.10)
    code.paragraph_format.right_indent = Inches(0.10)
    code._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")
    p = doc.add_paragraph(style="Title")
    set_fonts(p.add_run("SaddleLLM 产品手册与公众号宣传文案"), size=22)
    doc.add_paragraph("适用版本 2.40    2026 年 9 月 27 日")
    doc.add_paragraph("SaddleLLM 团队")
    doc.add_paragraph("本册包含可直接用于公众号的文章、方案对比与部署步骤，以及短宣传文案。当前最明确的优势是统一训练和服务入口，并提供自有原生优化及可选 vLLM 后端。性能数据保留实测条件与质量边界，便于读者判断自己的场景是否适用。")
    doc.add_paragraph("阅读顺序  先看公众号文章了解产品，再按对比与部署手册操作。短文案用于摘要、朋友圈和发布检查。")
    add_markdown(doc, SOURCE_DIR / "SaddleLLM_公众号文章.md")
    doc.add_page_break()
    add_markdown(doc, SOURCE_DIR / "SaddleLLM_推理部署对比与推广手册.md")
    doc.add_page_break()
    add_markdown(doc, SOURCE_DIR / "SaddleLLM_宣传短文案.md", compact=True)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.core_properties.title = "SaddleLLM 产品手册与公众号宣传文案"
    doc.core_properties.subject = "模型推理部署 对比与宣传"
    doc.core_properties.author = "SaddleLLM 团队"
    doc.core_properties.keywords = "SaddleLLM Qwen3 推理部署 W8A16 vLLM"
    # The bundled default Word template contains a blue Title paragraph rule.
    # Remove inherited/direct paragraph borders rather than hiding the symptom.
    for container in [doc.styles._element, doc._element]:
        for border in list(container.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
    doc.styles["Title"].font.underline = False
    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output)
    print(args.output)
    print(f"Paragraphs: {len(doc.paragraphs)}; tables: {len(doc.tables)}")


if __name__ == "__main__":
    main()
