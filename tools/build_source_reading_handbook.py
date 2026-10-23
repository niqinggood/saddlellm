"""Build the source-reading Word handbook with the Codex packaged Python."""
from __future__ import annotations

import json
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
QA = ROOT / 'outputs' / 'source_handbook_20261001'
OUTPUT = ROOT / 'SaddleLLM_主要功能源码解读与学习手册.docx'
MAIN = ROOT / 'docs' / 'CORE_FEATURE_CODE_WALKTHROUGH_ZH.md'
ROAD = ROOT / 'docs' / 'SOURCE_READING_ROADMAP_ZH.md'
QA.mkdir(parents=True, exist_ok=True)
FONT = 'Microsoft YaHei'
REFS: dict[str, tuple[str, str]] = {}
BOOKMARK_ID = 1


def font(run, size=None, bold=None, mono=False, color='000000'):
    run.font.name = 'Consolas' if mono else 'Calibri'
    run.font.color.rgb = RGBColor.from_string(color)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    rf = run._element.get_or_add_rPr().get_or_add_rFonts()
    rf.set(qn('w:eastAsia'), FONT)
    return run


def bookmark(p, name):
    global BOOKMARK_ID
    start, end = OxmlElement('w:bookmarkStart'), OxmlElement('w:bookmarkEnd')
    start.set(qn('w:id'), str(BOOKMARK_ID))
    end.set(qn('w:id'), str(BOOKMARK_ID))
    start.set(qn('w:name'), name)
    p._p.insert(0, start)
    p._p.append(end)
    BOOKMARK_ID += 1


def hyperlink(p, label, *, anchor=None, href=None, size=11):
    el = OxmlElement('w:hyperlink')
    if anchor:
        el.set(qn('w:anchor'), anchor)
    else:
        rid = p.part.relate_to(href, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
        el.set(qn('r:id'), rid)
    r = OxmlElement('w:r')
    prop = OxmlElement('w:rPr')
    fonts = OxmlElement('w:rFonts')
    for key, value in [('ascii', 'Calibri'), ('hAnsi', 'Calibri'), ('eastAsia', FONT)]:
        fonts.set(qn('w:' + key), value)
    prop.append(fonts)
    c = OxmlElement('w:color'); c.set(qn('w:val'), '244F72'); prop.append(c)
    sz = OxmlElement('w:sz'); sz.set(qn('w:val'), str(int(size * 2))); prop.append(sz)
    r.append(prop)
    t = OxmlElement('w:t'); t.text = label; r.append(t); el.append(r)
    p._p.append(el)


INLINE = re.compile(r'(\[[^\]]+\]\([^)]+\)|\*\*.+?\*\*|`[^`]+`)')


def inline(p, value, size=11, bold=False):
    for token in INLINE.split(value):
        if not token:
            continue
        match = re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)', token)
        if match:
            label, dest = match.groups()
            if '#L' in dest and dest.startswith('../'):
                if dest not in REFS:
                    REFS[dest] = (f'S{len(REFS)+1:02d}', label)
                ref = REFS[dest][0]
                font(p.add_run(label), size, bold)
                hyperlink(p, f'〔{ref}〕', anchor=ref, size=size)
            elif dest.endswith('SOURCE_READING_ROADMAP_ZH.md'):
                hyperlink(p, '本册阅读路线', anchor='reading_route', size=size)
            elif dest.endswith('CORE_FEATURE_CODE_WALKTHROUGH_ZH.md'):
                hyperlink(p, '本册源码解读', anchor='source_main', size=size)
            else:
                href = dest if '://' in dest else (ROOT / 'docs' / dest).resolve().as_uri()
                hyperlink(p, label, href=href, size=size)
        elif token.startswith('**'):
            font(p.add_run(token[2:-2]), size, True)
        elif token.startswith('`'):
            font(p.add_run(token[1:-1]), size, bold, mono=True)
        else:
            font(p.add_run(token), size, bold)


def figure_font(size, bold=False):
    return ImageFont.truetype('C:/Windows/Fonts/msyhbd.ttc' if bold else 'C:/Windows/Fonts/msyh.ttc', size)


def diagram(kind):
    if kind == 'system':
        im = Image.new('RGB', (1500, 820), 'white')
        dr = ImageDraw.Draw(im)
        boxes = {
            'entry': (35, 60, 300, 145, 'CLI / Python API'),
            'config': (385, 60, 675, 145, '配置解析与校验'),
            'plan': (760, 60, 1050, 145, '计划与 Stage 注册'),
            'orch': (1140, 60, 1465, 145, 'TrainingOrchestrator'),
            'data': (50, 280, 340, 365, '数据与 Tokenizer'),
            'train': (450, 280, 760, 365, '模型构建与训练'),
            'eval': (870, 280, 1110, 365, '评测与门禁'),
            'export': (1190, 280, 1450, 365, '模型导出'),
            'serve': (1180, 505, 1450, 590, '推理服务'),
            'existing': (830, 505, 1060, 590, '已有模型'),
            'world': (65, 660, 420, 755, '多模态与世界模型训练'),
            'spatial': (520, 660, 875, 755, '空间规划与反馈'),
            'state': (1050, 660, 1450, 755, '阶段状态与产物登记'),
        }
        def arrow(points):
            dr.line(points, fill='#6D8190', width=4)
            x,y=points[-1]; px,py=points[-2]
            import math
            a=math.atan2(y-py,x-px)
            dr.polygon([(x,y),(x-17*math.cos(a-.5),y-17*math.sin(a-.5)),(x-17*math.cos(a+.5),y-17*math.sin(a+.5))],fill='#6D8190')
        for pts in [[(300,103),(385,103)],[(675,103),(760,103)],[(1050,103),(1140,103)],
                    [(1300,145),(1300,215),(195,215),(195,280)],[(605,215),(605,280)],
                    [(340,323),(450,323)],[(760,323),(870,323)],[(1110,323),(1190,323)],
                    [(1320,365),(1320,505)],[(1060,548),(1180,548)],
                    [(195,215),(30,215),(30,615),(245,615),(245,660)],
                    [(420,708),(520,708)],[(1458,145),(1480,145),(1480,625),(1250,625),(1250,660)]]:
            arrow(pts)
        for x1,y1,x2,y2,label in boxes.values():
            dr.rounded_rectangle((x1,y1,x2,y2),radius=10,fill='#EFF4F7',outline='#7E94A4',width=2)
            fs=27 if len(label)>19 else 30
            dr.text(((x1+x2)/2,(y1+y2)/2),label,font=figure_font(fs),anchor='mm',fill='#111111')
    else:
        im=Image.new('RGB',(1500,1050),'white'); dr=ImageDraw.Draw(im)
        xs=[145,535,925,1320]
        labels=['客户端','API 路由','GenerationEngine','模型或原生会话']
        for x,label in zip(xs,labels):
            dr.rounded_rectangle((x-140,30,x+140,105),radius=8,fill='#EFF4F7',outline='#7E94A4',width=2)
            dr.text((x,67),label,font=figure_font(28),anchor='mm',fill='black')
            for y in range(120,1000,22): dr.line((x,y,x,y+11),fill='#A4B1BA',width=2)
        steps=[(0,1,190,'chat/completions 请求'),(1,1,290,'参数校验与准入'),(1,2,390,'准备请求或执行生成'),
               (2,2,490,'聊天模板与 tokenize'),(2,2,600,'上下文与输出预算检查'),(2,3,710,'prefill 和 decode'),
               (3,2,810,'token'),(2,1,890,'文本与计数'),(1,0,975,'JSON 或 SSE')]
        for a,b,y,label in steps:
            x1,x2=xs[a],xs[b]
            if a==b:
                dr.line([(x1,y-20),(x1+115,y-20),(x1+115,y+20),(x1,y+20)],fill='#526C80',width=3)
                dr.polygon([(x1,y+20),(x1+14,y+12),(x1+14,y+28)],fill='#526C80')
                dr.text((x1-20,y-47),label,font=figure_font(25),anchor='rm',fill='#111111')
            else:
                dr.line((x1,y,x2,y),fill='#526C80',width=3)
                s=1 if x2>x1 else -1
                dr.polygon([(x2,y),(x2-s*17,y-9),(x2-s*17,y+9)],fill='#526C80')
                dr.text(((x1+x2)/2,y-25),label,font=figure_font(25),anchor='mm',fill='#111111')
    path=QA/f'{kind}.png'; im.save(path); return path


def table(doc, rows, widths=None, small=False):
    n=len(rows[0]); t=doc.add_table(rows=1, cols=n)
    t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    if widths is None:
        widths={2:[1.65,4.88],3:[1.28,2.78,2.47],4:[1.35,1.48,1.9,1.8]}[n]
    for col,w in zip(t.columns,widths): col.width=Inches(w)
    borders=OxmlElement('w:tblBorders')
    for name in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+name)
        for key,value in [('val','single'),('sz','4'),('color','D9D9D9')]: el.set(qn('w:'+key),value)
        borders.append(el)
    t._tbl.tblPr.append(borders)
    for i,values in enumerate(rows):
        row=t.rows[0] if i==0 else t.add_row()
        if i==0:
            repeat=OxmlElement('w:tblHeader'); row._tr.get_or_add_trPr().append(repeat)
        no_split=OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(no_split)
        for j,(cell,value) in enumerate(zip(row.cells,values)):
            cell.width=Inches(widths[j]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr()
            shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'E2EBF2' if i==0 else ('F7F9FA' if i%2==0 else 'FFFFFF')); tcpr.append(shade)
            margins=OxmlElement('w:tcMar')
            for side,v in [('top',70),('bottom',70),('left',90),('right',90)]:
                item=OxmlElement('w:'+side); item.set(qn('w:w'),str(v)); item.set(qn('w:type'),'dxa'); margins.append(item)
            tcpr.append(margins)
            p=cell.paragraphs[0]; p.paragraph_format.line_spacing=1.13; p.paragraph_format.space_after=Pt(0)
            p.paragraph_format.keep_with_next=(i==0)
            if small and j==0: p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            inline(p,value,size=8.5 if small else 9.0,bold=i==0)
    spacer=doc.add_paragraph(); spacer.paragraph_format.space_after=Pt(3); spacer.paragraph_format.space_before=Pt(0)
    spacer.paragraph_format.line_spacing=1; font(spacer.add_run(''),3)
    return t


def field(p,instruction,text=''):
    f=OxmlElement('w:fldSimple'); f.set(qn('w:instr'),instruction)
    if text:
        r=OxmlElement('w:r'); tx=OxmlElement('w:t'); tx.text=text; r.append(tx); f.append(r)
    p._p.append(f)


def style_doc(doc):
    sec=doc.sections[0]; sec.page_width=Inches(8.5); sec.page_height=Inches(11)
    sec.left_margin=sec.right_margin=Cm(2.2); sec.top_margin=Cm(2); sec.bottom_margin=Cm(1.9)
    sec.header_distance=Cm(.8); sec.footer_distance=Cm(.85)
    sec.different_first_page_header_footer=True
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','List Bullet','List Number','Caption']:
        st=doc.styles[name]; st.font.name='Calibri'; st.font.color.rgb=RGBColor(0,0,0)
        st._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),FONT)
        st.paragraph_format.widow_control=True
    st=doc.styles['Normal']; st.font.size=Pt(11); st.paragraph_format.line_spacing=1.2; st.paragraph_format.space_after=Pt(6)
    for name,size in [('Title',25),('Heading 1',16),('Heading 2',12),('Heading 3',11)]:
        st=doc.styles[name]; st.font.size=Pt(size); st.paragraph_format.keep_with_next=True
        st.paragraph_format.space_before=Pt(13 if name!='Title' else 0); st.paragraph_format.space_after=Pt(7)
    for name in ['Code','Flow','SourceIndex']:
        st=doc.styles.add_style(name,1); st.base_style=doc.styles['Normal']
        st.font.name='Consolas' if name=='Code' else 'Calibri'; st.font.size=Pt(9)
        st._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),FONT)
        st.paragraph_format.line_spacing=1.08; st.paragraph_format.space_after=Pt(1)
        st.paragraph_format.left_indent=Cm(.25); st.paragraph_format.right_indent=Cm(.2)
    header=sec.header.paragraphs[0]; header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    font(header.add_run('SaddleLLM 主要功能源码解读与学习手册'),8)
    footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    font(footer.add_run('第 '),8); field(footer,'PAGE','1'); font(footer.add_run(' 页'),8)
    settings=doc.settings._element
    update=OxmlElement('w:updateFields'); update.set(qn('w:val'),'true'); settings.append(update)


def markdown(doc,path,route=False):
    lines=path.read_text(encoding='utf-8').splitlines()
    start=next(i for i,l in enumerate(lines) if l.startswith('## '))
    i=start
    while i<len(lines):
        raw=lines[i]; line=raw.strip()
        if not line: i+=1; continue
        if line.startswith('~~~'):
            lang=line[3:]; body=[]; i+=1
            while i<len(lines) and not lines[i].startswith('~~~'):
                body.append(lines[i]); i+=1
            if lang=='mermaid':
                kind='sequence' if body[0].startswith('sequenceDiagram') else 'system'
                p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(str(diagram(kind)),width=Cm(16.1))
                p.paragraph_format.keep_with_next=True
                cap=doc.add_paragraph('图 2 聊天请求时序' if kind=='sequence' else '图 1 系统主流程',style='Caption')
                cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
            else:
                for k,value in enumerate(body):
                    p=doc.add_paragraph(style='Code' if lang in ['python','powershell','yaml'] else 'Flow')
                    font(p.add_run(value or ' '),9,mono=lang in ['python','powershell','yaml'])
                    p.paragraph_format.keep_with_next=(len(body)<=12 and k<len(body)-1)
                    p.paragraph_format.keep_together=True
                    if k==len(body)-1: p.paragraph_format.space_after=Pt(8)
            i+=1; continue
        if line.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                cells=[c.strip() for c in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?',c) for c in cells): rows.append(cells)
                i+=1
            if len(rows[0])==4 and rows[0][0]=='阅读顺序':
                widths=[1.12,1.43,1.92,2.06]
            elif rows[0][0]=='对象': widths=[1.33,2.23,2.97]
            elif rows[0][0]=='现象': widths=[1.65,2.3,2.58]
            else: widths=None
            table(doc,rows,widths); continue
        if line.startswith('## '):
            title=line[3:]
            p=doc.add_heading(title,level=2 if route else 1)
            if not route and title.startswith('1 '): bookmark(p,'source_main')
            i+=1; continue
        if line.startswith('- '):
            p=doc.add_paragraph(style='List Bullet'); inline(p,line[2:])
        elif re.match(r'^\d+\. ',line):
            p=doc.add_paragraph(style='List Number'); inline(p,re.sub(r'^\d+\. ','',line))
        else:
            p=doc.add_paragraph(); inline(p,line)
        i+=1


def main():
    doc=Document(); style_doc(doc)
    doc.add_paragraph('SaddleLLM',style='Title')
    doc.add_paragraph('主要功能源码解读与学习手册',style='Title')
    doc.add_paragraph('源码基线 2.40    2026 年 10 月 1 日',style='Subtitle')
    doc.add_paragraph('这份手册将主要功能的实现流程与 12 讲阅读路线整理在一起，帮助你从一个命令追踪到关键函数、数据变化和最终产物，逐步建立独立维护系统的能力。')
    doc.add_paragraph('内容基于当前工作区源码，包含本地未提交实现。先掌握配置、Stage 编排和状态恢复，再进入数据、模型、训练与推理，最后阅读世界模型和空间应用。')
    p=doc.add_paragraph(); font(p.add_run('如何使用这份手册'),12,True)
    for value in ['第一次建议用约 60 至 90 分钟完成阅读路线的前三讲。','每条功能链配有源码位置、关键代码、断点建议、关联测试与自查题。','正文中的 S 编号可跳到文末源码索引；索引保留项目相对路径和行号，便于在 IDE 中定位。']:
        p=doc.add_paragraph(style='List Bullet'); inline(p,value)
    doc.add_paragraph('源码行号与配置能力以本次核对为准。运行计划和轻量测试验证流程契约，训练质量与 GPU 性能仍需在具体环境中验证。')
    p=doc.add_paragraph(); font(p.add_run('建议首先追踪'),12,True)
    for value in ['train 命令 → 配置解析 → PipelinePlan','TrainingOrchestrator.run → run_stage → 执行器','结果字典 → 产物登记 → 状态保存与恢复']:
        doc.add_paragraph(value,style='Flow')
    doc.add_page_break()
    p=doc.add_paragraph(); font(p.add_run('目录'),20,True)
    p=doc.add_paragraph(); field(p,'TOC \\o "1-1" \\h \\z \\u')
    doc.add_page_break()
    p=doc.add_heading('阅读路线与练习',level=1); bookmark(p,'reading_route')
    markdown(doc,ROAD,route=True)
    doc.add_page_break()
    markdown(doc,MAIN)
    doc.add_page_break()
    doc.add_heading('关键源码索引',level=1)
    doc.add_paragraph('位置均相对项目根目录。L 表示当前源码行号；后续代码变动时优先按函数名搜索。正文中的 S 编号链接到对应记录。')
    refs=list(REFS.items())
    rows=[['编号','类或函数','项目相对路径与行号']]
    for dest,(ref,label) in refs:
        rel,line=dest[3:].rsplit('#L',1)
        parts=rel.rsplit('/',1)
        location=(parts[0]+'/\n' if len(parts)>1 else '')+parts[-1]+'\nL'+line
        rows.append([ref,label,location])
    t=table(doc,rows,[.5,2.75,3.28],small=True)
    for row,(_, (ref,_)) in zip(t.rows[1:],refs): bookmark(row.cells[0].paragraphs[0],ref)
    for container in [doc.styles._element,doc._element]:
        for el in list(container.iter(qn('w:pBdr'))): el.getparent().remove(el)
    doc.styles['Title'].font.underline=False
    doc.core_properties.title='SaddleLLM 主要功能源码解读与学习手册'
    doc.core_properties.subject='主要功能调用链 源码阅读 学习路线'
    doc.core_properties.author=''
    doc.core_properties.keywords='SaddleLLM 源码 训练 推理 世界模型'
    doc.save(OUTPUT)
    (QA/'build_summary.json').write_text(json.dumps({'output':str(OUTPUT),'source_symbols':len(refs),'tables':len(doc.tables),'paragraphs':len(doc.paragraphs)},ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'output':str(OUTPUT),'source_symbols':len(refs),'tables':len(doc.tables),'paragraphs':len(doc.paragraphs)},ensure_ascii=False))


if __name__=='__main__': main()
